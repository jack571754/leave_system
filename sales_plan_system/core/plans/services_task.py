import csv
import io
from decimal import Decimal, InvalidOperation
from typing import List, Optional

from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from core.approvals.models import TaskApprovalInstance, TaskApprovalRecord
from core.errors import ValidationError
from core.products.models import ProductSKU

from .models import PlanCampaign, PlanTask, PlanTaskItem
from .service_results import TaskResult


def _parse_decimal(value, *, field: str) -> Decimal:
    try:
        return Decimal(str(value if value not in (None, "") else 0))
    except InvalidOperation as exc:
        raise ValidationError([{"field": field, "message": "必须是合法数字"}]) from exc


def _parse_int(value, *, field: str) -> int:
    try:
        result = int(value if value not in (None, "") else 0)
    except (TypeError, ValueError) as exc:
        raise ValidationError([{"field": field, "message": "必须是整数"}]) from exc
    if result < 0:
        raise ValidationError([{"field": field, "message": "不能为负数"}])
    return result


class TaskService:
    PREVIEW_SESSION_KEY = "task_import_preview"

    def create_campaign(
        self,
        *,
        created_by,
        title: str,
        plan_month: str,
        plan_type: str,
        start_date,
        end_date,
        org_unit=None,
        stores=None,
        remark: str = "",
    ):
        if not title or not plan_month:
            raise ValidationError([{"field": "campaign", "message": "请填写任务标题和计划月份"}])
        if start_date and end_date and start_date > end_date:
            raise ValidationError([{"field": "date", "message": "开始日期不能晚于结束日期"}])

        stores = list(stores or [])
        if not stores:
            raise ValidationError([{"field": "stores", "message": "请至少选择一个店铺"}])

        with transaction.atomic():
            campaign = PlanCampaign.objects.create(
                title=title,
                plan_month=plan_month,
                plan_type=plan_type,
                start_date=start_date,
                end_date=end_date,
                org_unit=org_unit,
                created_by=created_by,
                status=PlanCampaign.Status.ACTIVE,
                remark=remark,
                published_at=timezone.now(),
            )
            for store in stores:
                PlanTask.objects.create(
                    campaign=campaign,
                    store=store,
                    org_unit=store.org_unit,
                    assignee=store.owner,
                    editable_until=(
                        timezone.make_aware(
                            timezone.datetime.combine(
                                end_date,
                                timezone.datetime.max.time().replace(microsecond=0),
                            )
                        )
                        if end_date
                        else None
                    ),
                )
        return campaign

    def add_task_items(
        self,
        task: PlanTask,
        sku_ids: List[int],
        unit_price: Optional[Decimal] = None,
        remark: str = "",
    ) -> TaskResult:
        if not sku_ids:
            raise ValidationError([{"field": "sku_ids", "message": "请先选择要加入任务的 SKU"}])

        skus = ProductSKU.objects.filter(pk__in=sku_ids, status=ProductSKU.Status.ACTIVE)
        target_months = task.get_target_months()
        created_count = 0

        for sku in skus:
            price = unit_price if unit_price is not None else sku.unit_price
            for target_month in target_months:
                PlanTaskItem.objects.get_or_create(
                    task=task,
                    sku=sku,
                    target_month=target_month,
                    defaults={
                        "sales_qty": 0,
                        "unit_price": price,
                        "stock_qty": 0,
                        "remark": remark,
                    },
                )
            created_count += 1

        return TaskResult(success=True, task=task, message=f"已添加 {created_count} 个 SKU")

    def save_task_matrix(self, task: PlanTask, matrix_data: dict) -> TaskResult:
        target_months = task.get_target_months()
        sku_ids = matrix_data.get("sku_ids", [])
        updated_rows = 0

        with transaction.atomic():
            for sku_id in sku_ids:
                unit_price = _parse_decimal(matrix_data.get(f"unit_price__{sku_id}"), field=f"unit_price__{sku_id}")
                for month in target_months:
                    qty_key = f"sales_qty__{sku_id}__{month}"
                    stock_key = f"stock_qty__{sku_id}__{month}"
                    sales_qty = _parse_int(matrix_data.get(qty_key), field=qty_key)
                    stock_qty = _parse_int(matrix_data.get(stock_key), field=stock_key)

                    item = task.items.filter(sku_id=sku_id, target_month=month).first()
                    if not item:
                        continue

                    item.sales_qty = sales_qty
                    item.stock_qty = stock_qty
                    item.unit_price = unit_price
                    item.save(update_fields=["sales_qty", "stock_qty", "unit_price", "updated_at"])

                updated_rows += 1

            deleted_ids = matrix_data.get("delete_sku_ids", [])
            if deleted_ids:
                task.items.filter(sku_id__in=deleted_ids).delete()

        return TaskResult(success=True, task=task, message=f"已更新 {updated_rows} 行数据")

    def bulk_delete_items(self, task: PlanTask, sku_ids: List[int]) -> TaskResult:
        if not sku_ids:
            raise ValidationError([{"field": "sku_ids", "message": "请至少选择一个 SKU"}])
        deleted_count, _ = task.items.filter(sku_id__in=sku_ids).delete()
        return TaskResult(success=True, task=task, message=f"已删除 {deleted_count} 条明细")

    def build_import_template(self, task: PlanTask):
        workbook = Workbook()
        sheet = workbook.active
        headers = ["sku_code", "unit_price"]
        for month in task.get_target_months():
            headers.extend([f"{month}_sales_qty", f"{month}_stock_qty"])
        sheet.append(headers)
        return workbook

    def parse_import_file(self, task: PlanTask, uploaded_file):
        target_months = task.get_target_months()
        rows = []
        errors = []

        if uploaded_file.name.lower().endswith(".csv"):
            text = uploaded_file.read().decode("utf-8-sig")
            raw_rows = list(csv.DictReader(io.StringIO(text)))
        else:
            workbook = load_workbook(uploaded_file, data_only=True)
            sheet = workbook.active
            header_row = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            raw_rows = [dict(zip(header_row, values)) for values in sheet.iter_rows(min_row=2, values_only=True)]

        for index, raw_row in enumerate(raw_rows, start=2):
            sku_code = str(raw_row.get("sku_code") or "").strip()
            if not sku_code:
                errors.append({"row": index, "message": "sku_code 不能为空"})
                continue

            sku = ProductSKU.objects.filter(sku_code=sku_code, status=ProductSKU.Status.ACTIVE).first()
            if not sku:
                errors.append({"row": index, "message": f"SKU {sku_code} 不存在或未启用"})
                continue

            try:
                unit_price = Decimal(str(raw_row.get("unit_price") or sku.unit_price))
            except InvalidOperation:
                errors.append({"row": index, "message": f"SKU {sku_code} 的单价格式不正确"})
                continue

            payload = {
                "sku_id": sku.pk,
                "sku_code": sku.sku_code,
                "sku_name": sku.sku_name,
                "unit_price": str(unit_price),
                "months": {},
            }
            valid = True
            for month in target_months:
                sales_key = f"{month}_sales_qty"
                stock_key = f"{month}_stock_qty"
                try:
                    sales_qty = int(raw_row.get(sales_key) or 0)
                    stock_qty = int(raw_row.get(stock_key) or 0)
                except (TypeError, ValueError):
                    errors.append({"row": index, "message": f"{sku_code} 的销量或库存格式不正确"})
                    valid = False
                    break
                if sales_qty < 0 or stock_qty < 0:
                    errors.append({"row": index, "message": f"{sku_code} 的销量或库存不能为负数"})
                    valid = False
                    break
                payload["months"][month] = {"sales_qty": sales_qty, "stock_qty": stock_qty}

            if valid:
                rows.append(payload)

        return rows, errors

    def confirm_import_rows(self, task: PlanTask, preview_rows: List[dict]) -> TaskResult:
        with transaction.atomic():
            for row in preview_rows:
                for month, values in row["months"].items():
                    PlanTaskItem.objects.update_or_create(
                        task=task,
                        sku_id=row["sku_id"],
                        target_month=month,
                        defaults={
                            "sales_qty": values["sales_qty"],
                            "stock_qty": values["stock_qty"],
                            "unit_price": Decimal(row["unit_price"]),
                        },
                    )
        return TaskResult(success=True, task=task, message=f"已导入 {len(preview_rows)} 个 SKU")

    def submit_task(self, task: PlanTask, user, comment: str = "") -> TaskResult:
        if not task.can_submit(user):
            raise ValidationError([{"field": "task", "message": "当前任务不可提交，请检查状态、时间窗口或明细数据"}])

        steps = task.get_approval_steps()
        if not steps:
            raise ValidationError([{"field": "approval_flow", "message": "未找到可用的审批链"}])

        missing_labels = [step["label"] for step in steps if not step.get("approver")]
        if missing_labels:
            raise ValidationError(
                [{"field": "approval_flow", "message": f"以下审批步骤缺少审批人: {', '.join(missing_labels)}"}]
            )

        steps_snapshot = [
            {
                "step": step["step"],
                "approver_id": step["approver"].pk,
                "approver_name": step["approver"].get_full_name() or step["approver"].username,
                "label": step["label"],
            }
            for step in steps
        ]

        with transaction.atomic():
            next_round = (task.round_no or 0) + 1
            instance = TaskApprovalInstance.objects.create(
                task=task,
                round_no=next_round,
                status=TaskApprovalInstance.Status.PENDING,
                current_step=1,
                submitter=user,
                pending_approver=steps[0]["approver"],
                submit_comment=comment,
                steps_snapshot=steps_snapshot,
            )
            TaskApprovalRecord.objects.create(
                instance=instance,
                task=task,
                approver=user,
                approval_level=0,
                action=TaskApprovalRecord.Action.RESUBMIT,
                comment=comment,
            )

            now = timezone.now()
            task.submitter = user
            task.submitted_at = task.submitted_at or now
            task.last_submitted_at = now
            task.sync_from_instance(instance)
            task.save(
                update_fields=[
                    "submitter",
                    "submitted_at",
                    "last_submitted_at",
                    "status",
                    "approval_status",
                    "current_step",
                    "pending_approver",
                    "round_no",
                    "updated_at",
                ]
            )

        return TaskResult(success=True, task=task, message=f"{task.store.name} 任务已提交审批")

    def withdraw_task(self, task: PlanTask, user, comment: str = "") -> TaskResult:
        if not task.can_withdraw(user):
            raise ValidationError([{"field": "task", "message": "当前任务不能撤回"}])

        instance = task.get_active_approval_instance()
        if not instance or instance.status != TaskApprovalInstance.Status.PENDING:
            raise ValidationError([{"field": "task", "message": "没有可撤回的审批实例"}])

        with transaction.atomic():
            now = timezone.now()
            instance.status = TaskApprovalInstance.Status.WITHDRAWN
            instance.pending_approver = None
            instance.withdrawn_at = now
            instance.finished_at = now
            instance.save(update_fields=["status", "pending_approver", "withdrawn_at", "finished_at", "updated_at"])

            TaskApprovalRecord.objects.create(
                instance=instance,
                task=task,
                approver=user,
                approval_level=instance.current_step,
                action=TaskApprovalRecord.Action.WITHDRAW,
                comment=comment,
            )

            task.last_withdrawn_at = now
            task.sync_from_instance(instance)
            task.save(
                update_fields=[
                    "status",
                    "approval_status",
                    "current_step",
                    "pending_approver",
                    "last_withdrawn_at",
                    "updated_at",
                ]
            )

        return TaskResult(success=True, task=task, message=f"{task.store.name} 任务已撤回")
